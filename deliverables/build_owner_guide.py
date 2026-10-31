"""Build a beginner Word handbook and fictional test files; does not alter app logic or hosted data."""
from pathlib import Path
from decimal import Decimal as D
from datetime import date
import os, sys, csv, io, zipfile, re
sys.path.insert(0, str(Path(os.environ['TEMP']) / 'profitplate-doc-guide'))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'web' / 'public' / 'guides'
OUT.mkdir(parents=True, exist_ok=True)
NAME = 'ProfitPlate_Kitchen_Owner_Beginner_Guide'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.72)
sec.header_distance = sec.footer_distance = Inches(.30)
sec.different_first_page_header_footer = True
styles = doc.styles
for name in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    styles[name].font.name = 'Calibri'
    styles[name].font.size = Pt(10.5)
    styles[name].paragraph_format.space_after = Pt(5)
    styles[name].paragraph_format.line_spacing = 1.08
for name, size in [('Title', 32), ('Heading 1', 20), ('Heading 2', 13), ('Heading 3', 11)]:
    styles[name].font.name = 'Calibri'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string('173E32')
    styles[name].paragraph_format.keep_with_next = True
styles['Normal'].font.color.rgb = RGBColor.from_string('243D35')
header = sec.header.paragraphs[0]
header.text = 'PROFITPLATE  |  OWNER SETUP & TEST HANDBOOK'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('597365')
footer = sec.footer.paragraphs[0]
footer.text = 'Test restaurant only  •  Guide v1.0  |  Page '
footer.runs[0].font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

UI_NAMES = ['Team and access','Additional operational duties','Assign additional duty','Generate invitation code','Business access code','Join business','Manager appointments','Appoint manager','Operating controls','Set up the three areas','Save operating areas','Create open period','Reopen period','Add inventory item','Save supplier item','Save vendor','Save menu mapping','Save recipe','Add another item','Save purchase order draft','Submit order for approval','Approve purchase order','Record accepted delivery','Save receiving draft','Submit for approval','Approve and post','Save request draft','Submit request','Approve request','Dispatch approved stock','Dispatch entered quantities','Confirm received quantities','Confirm receipt','Validate and save sales draft','View report','Submit reconciliation','Approve reconciliation','Publish snapshot','View saved report','Download spreadsheet (CSV)','Save count draft','Submit count for approval','Approve count adjustment','Save action progress','Create action','Create workspace']
UI_PATTERN = re.compile('(' + '|'.join(re.escape(x) for x in sorted(UI_NAMES,key=len,reverse=True)) + ')')

def p(text='', bold=False, style=None):
    para = doc.add_paragraph(style=style)
    if bold:
        para.add_run(text).bold=True
    else:
        for part in UI_PATTERN.split(text):
            if part: para.add_run(part).bold = part in UI_NAMES
    return para

def note(text):
    para = p(text)
    props = para._p.get_or_add_pPr()
    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F2F0E5'); props.append(shade)
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(9)
    return para

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, label in enumerate(headers): t.rows[0].cells[i].text = str(label)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for cell in t.rows[0].cells:
        sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '173E32'); cell._tc.get_or_add_tcPr().append(sh)
        for r in cell.paragraphs[0].runs: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    for row in t.rows:
        trp = row._tr.get_or_add_trPr(); trp.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.0
                for r in para.runs: r.font.size = Pt(9)
    p()
    return t

def steps(items):
    for i, text in enumerate(items, 1): p(f'{i}. {text}')

def bullets(items):
    for text in items: p(text, style='List Bullet')

def chapter(number, title, account=None, route=None):
    doc.add_page_break()
    doc.add_heading(f'{number}. {title}', 1)
    if account: p('USE THIS ACCOUNT: ' + account, True)
    if route: p('OPEN: ' + route, True)

def check(text): p('[ ] PASS CHECK: ' + text, True)
def sub(text): doc.add_heading(text, 2)

doc.core_properties.title = 'ProfitPlate Kitchen — Owner Beginner Setup and Testing Guide'
doc.core_properties.subject = 'Step-by-step owner setup, master data, roles and fictional restaurant acceptance tests'
doc.core_properties.author = 'ProfitPlate'
doc.core_properties.keywords = 'beginner, owner, restaurant, setup, testing, inventory, reports'
p('PROFITPLATE', True)
doc.add_heading('Your first restaurant\nsetup & testing guide', 0)
p('For the Owner of ProfitPlate Kitchen', True)
p('A beginner-friendly Word handbook with exact screen labels, practical examples and expected results.')
p('Version 1.0  •  27 October 2026')
note('This is a fictional test exercise for the current application. It does not add users or stock for you. Follow the steps in the test restaurant only. Your existing 327 catalogue items should remain intact.')
sub('What you will accomplish')
bullets(['Set up test accounts and assign roles without using real staff accounts.',
         'Check Main Store, Kitchen and Bar, and open a valid inventory period.',
         'Add five clearly marked test SKUs, a vendor, supplier packs and recipes.',
         'Run purchasing, receiving, department requests, sales, production, waste and counts.',
         'Check COGS, gross margin, food cost, published reports and spreadsheet exports.',
         'Record pass/fail results and recognise safe rejection messages.'])
sub('Keep these details beside you')
table(['Detail', 'Write your value here'], [
    ('App address used on your laptop', '________________________________________'),
    ('Owner email (never write the password here)', '________________________________________'),
    ('Test date D — yesterday / a completed business date', '________________________________________'),
    ('Later test date T — today', '________________________________________'),
    ('Date testing started', '________________________________________')])
p('Suggested pace: complete Setup first, then the Baseline Test, then Advanced Tests. Stop at each PASS CHECK before continuing.')

chapter('START', 'How to use this handbook')
sub('Read the account label before every exercise')
p('You are the Owner. For testing you will also give your Owner account explicit Procurement and Inventory duties. A different Storekeeper account receives your purchase orders, and a different Operations Manager account approves transactions. Kitchen and Bar accounts confirm their own deliveries.')
note('One person may simulate the test accounts, but they are separate identities in the system. Using your Owner login on a second phone does not create an independent approver. Real live approvals should be carried out by the designated people.')
sub('Choose two dates before you start')
steps(['Set D to yesterday, using the Africa/Lagos business date. Choose a date with no other posted sales in this test restaurant. The baseline report assumes exactly the two sample sales lines in Appendix A.',
       'Set T to today. Use D for the initial purchase receipts, department dispatches and baseline sales. Use T for later cost-change, count, waste, usage and production exercises.',
       'This backdating is only to test fictional data and the completed-day reconciliation screen immediately. Real transactions should use their true business dates.',
       'Check Reports for D first. If other posted sales or unrelated pending stock records already exist on D, select another completed test date or ask for help arranging an isolated test workspace. Do not delete existing business history to force the example totals.'])
sub('Find your way around')
p('In Microsoft Word, use View → Navigation Pane → Headings to jump between sections. Bold button names in the instructions match controls in the app.')
table(['On a laptop', 'On a phone'], [('Use the left navigation menu.', 'Tap Menu at the top to open navigation.'),
    ('Scroll down: forms may be below existing records.', 'Swipe a wide table horizontally inside its panel.'),
    ('Open a small arrow/summary to expand actions.', 'Tap the same text to expand a collapsed action.')])
sub('Suggested reading order')
table(['Sections', 'Purpose'], [('1–5', 'Accounts, roles, locations and periods'), ('6–9', 'Items, supplier packs, recipes and menu mappings'),
    ('10–14', 'Baseline: buy, receive, distribute, import sales and prove the report'), ('15–20', 'Cost changes, counts, waste, production and permission tests'),
    ('21–23', 'Troubleshooting, test record and glossary'), ('Appendices A–C', 'CSV examples, arithmetic and completion checklist')])

chapter(1, 'Open the app and confirm your Owner account', 'OWNER')
steps(['Keep the current development server running. On your laptop, open the working app address. If you normally use http://localhost:3000, use that address.',
       'If the server is stopped and you normally start it yourself: open a normal PowerShell window, run cd C:\\Users\\Hp\\Desktop\\ProfitPlate\\web, then run npm.cmd run dev. Leave that terminal open while testing.',
       'Sign in with your existing Owner email and password. If you have already created the restaurant, continue in that workspace; do not create another one.',
       'Only if this is a brand-new Owner account with no restaurant: choose Create an account, complete email verification in that same browser profile, and enter ProfitPlate Kitchen in Business name. Click Create workspace.',
       'Confirm the workspace header identifies your test restaurant. Open Business settings and confirm your role is owner, currency is NGN and timezone is Africa/Lagos.',
       'If a business switcher appears, choose the intended test restaurant before entering any records.'])
note('The current account setup allows one owned restaurant per user. A restaurant Administrator is not the overall ProfitPlate platform administrator. This handbook uses your restaurant Owner account.')
check('You can sign out, sign in again and return to the correct restaurant. Refreshing does not remove the business.')
sub('Phone access')
p('Use the app address already working on your phone. localhost on your phone points to the phone itself, so do not copy the laptop’s localhost address to the phone. If the network changes and the link stops working, resolve the address/server connection first; changing inventory settings will not fix connectivity.')

chapter(2, 'Prepare your test accounts', 'OWNER')
p('Use email addresses you control and can verify. The names below are labels, not ready-made email accounts. No real employee accounts are required. Keep passwords in a password manager, not in this guide.')
p('If your email provider already supports aliases or plus-addresses, distinct verified test addresses may deliver to one inbox. Each must still be a separate ProfitPlate login. Otherwise use separate test inboxes; made-up addresses cannot complete email verification.')
table(['Account label', 'Working role / duty', 'Purpose'], [
 ('OWNER', 'Existing Owner; add Procurement and Inventory duties in section 4', 'Masters, purchase orders, sales import and Finance test'),
 ('OPS', 'Invite as Viewer; separately appoint Operations Manager', 'Independent approval and period control'),
 ('STORE', 'Storekeeper / Inventory Manager', 'Receive orders and dispatch stock'),
 ('KITCHEN', 'Kitchen Lead', 'Kitchen requests, receipts, waste and production'),
 ('BAR', 'Bar Lead', 'Bar requests, receipts and waste'),
 ('FINANCE — optional', 'Accountant', 'Separate Finance workspace and reconciliation test'),
 ('ADMIN — optional', 'Administrator', 'Check setup access without automatic financial/posting access'),
 ('OPS BACKUP — optional', 'Viewer plus Operations Manager appointment', 'Backup approval and manager-created transaction tests')])
sub('Keep sessions separate')
steps(['Create separate normal browser profiles named OWNER, OPS, STORE, KITCHEN and BAR, or sign out before changing accounts.',
       'Keep the Owner profile open for invitations and appointments. Open the staff profile separately for sign-up and joining.',
       'Always check the signed-in email before clicking Submit or Approve. Multiple private/incognito windows may share one private session; separate normal profiles are clearer.'])
note('For the main exercise, OPS only approves records created by other accounts. Add a second Operations Manager before testing approval of records created by OPS. Nobody receives a self-approval exception.')
check('You have selected controlled, verifiable test emails and know which browser profile belongs to each account.')

chapter(3, 'Invite each account into your restaurant', 'OWNER, then each INVITED ACCOUNT', 'Team and access → Invite a team member')
steps(['In the Owner profile, open Team and access. Scroll to Invite a team member.',
       'Enter the test email in Invited email. Choose its Working role using the table in section 2. For OPS choose Viewer first.',
       'Click Generate invitation code. Copy the full Invitation access code privately. This app flow does not send an invitation email.',
       'Open the intended person’s separate browser profile. If that email has no ProfitPlate account, create one and complete its email verification. If it already has an account, sign in.',
       'On the business setup screen, use Enter your team access code. If already in another workspace, choose Join another business. You can also open the same app address followed by /join.',
       'Confirm the Join screen shows the invited email. Paste the full code into Business access code and click Join business.',
       'Return to the Owner profile, refresh Team and access and check the person appears as active. Repeat for STORE, KITCHEN, BAR and OPS.'])
sub('If an invitation fails')
bullets(['The signed-in email must match the invited email exactly. Verification must be completed.',
         'Copy the whole code, including its dot-separated parts. Do not place it in the page address or a screenshot.',
         'If a code expires or is lost, open Invitations → Replace or revoke access code → Replace access code. The old code then stops working.',
         'Do not create a new restaurant for each staff account. They should join your existing restaurant.'])
check('All four invited accounts appear as active members of ProfitPlate Kitchen.')

chapter(4, 'Assign your testing duties and appoint Operations', 'OWNER', 'Team and access')
sub('Give your Owner account two explicit duties')
steps(['Scroll to Additional operational duties. Select your Owner email as Duty holder.',
       'Select Procurement Officer under Additional duty. Enter Owner test walkthrough in Reason / reference. Click Assign additional duty.',
       'Repeat with your Owner email and Inventory / Storekeeper. Use the same clear reason and click Assign additional duty.',
       'Refresh the workspace. Purchasing creation and sales-import actions should now be available to your Owner login. Your Owner role has not been replaced.'])
sub('Appoint the independent approver')
steps(['Still as Owner, scroll to Manager appointments.',
       'Select the active OPS test email in Active team member.',
       'Select Operations Manager as Manager role. Click Appoint manager.',
       'In the OPS profile, refresh the workspace. It should show the Operations Manager workspace and Approvals navigation.'])
note('Do not appoint your Owner login as the only approver for this walkthrough. Owner creates the test purchase order and sales import; OPS must be a different account. STORE must receive the order because its creator cannot receive it.')
check('Owner can prepare purchases and sales imports. OPS can review approvals. STORE remains a separate receiver.')
sub('Optional separate Finance and Administrator checks')
p('Invite FINANCE as Accountant and ADMIN as Administrator using section 3. Finance can prepare reconciliations but cannot approve stock adjustments unless separately appointed Operations Manager. Administrator has setup access and does not automatically receive posting or financial-report rights.')

chapter(5, 'Check Main Store, Kitchen, Bar and your open period', 'OWNER, then OPS', 'Inventory → Operating controls')
sub('Owner: configure the operating areas')
steps(['Open Inventory → Operating controls. Find Three operating areas.',
       'If Main Store, Kitchen and Bar are already mapped correctly, keep the mappings and continue.',
       'If setup is incomplete, click Set up the three areas. The app tries to reuse matching existing locations, including Kitchen.',
       'If it asks you to choose manually, expand Choose existing locations manually. Select three different locations in Main Store, Kitchen and Bar. Click Save operating areas.',
       'Confirm all three area names appear. Do not remap them after stock movements have started.'])
sub('OPS: make sure a valid open inventory period exists')
steps(['In the OPS profile, open Inventory → Operating controls → Inventory periods.',
       'Check that D and T each fall within an existing Open period. Two adjacent open periods can cover the dates separately. If both dates are covered, use those periods; do not add an overlapping period.',
       'If no suitable period exists, enter Start date and End date covering your chosen test dates and click Create open period. Choose a non-overlapping range.',
       'If a required existing period is Closed, use Reopen period with an appropriate test reason. Only Operations Manager can perform this action.'])
check('The three locations are distinct, and both D and T are within an Open period before any stock posting.')

chapter(6, 'Add the five test inventory items', 'OWNER with Inventory duty', 'Master data → Items and categories, or Inventory → Items')
p('Keep the imported catalogue intact. Create the following five new, clearly marked items. If any GUIDE code already exists from a previous attempt, inspect it before continuing; do not create a duplicate or reuse an item with existing test movements.')
table(['SKU', 'Item name', 'Category', 'Base unit'], [
 ('GUIDE-RICE', 'Guide Rice', 'Raw material', 'kg'),
 ('GUIDE-FLOUR', 'Guide Flour', 'Raw material', 'kg'),
 ('GUIDE-SYRUP', 'Guide Syrup', 'Raw material', 'l'),
 ('GUIDE-SAUCE', 'Guide Sauce', 'Semi-finished', 'kg'),
 ('GUIDE-TISSUE', 'Guide Tissue', 'General supply', 'each')])
steps(['Scroll to Add inventory item.',
       'Enter Item name and SKU from the first row. Choose Category and Base unit exactly as shown.',
       'Click Add inventory item. Find the saved row in the catalogue.',
       'Repeat for the other four rows. Refresh and confirm all five remain.'])
note('Base unit is the unit in which stock is counted. A rice bag is a purchase pack; rice stock is counted in kg. The semi-finished Guide Sauce is produced in Kitchen and must not be bought from the supplier in this exercise.')
check('Five GUIDE items exist with the correct categories and units. They have no stock merely because they have been created.')

chapter(7, 'Add a vendor and verified purchase packs', 'OWNER with Procurement duty', 'Master data → Vendors, then Supplier items')
sub('Add the vendor')
steps(['Open Master data → Vendors. Scroll to Add vendor.',
       'Enter Vendor code GUIDE-VENDOR and Vendor name Guide Test Supplier.',
       'Leave optional contact/email/phone/address fields blank, or use controlled fictional test details. Keep Status Active.',
       'Click Save vendor. Confirm the vendor row appears.'])
sub('Link each item to its purchase pack')
table(['Item', 'Purchase unit / pack', 'Base units per purchase unit', 'Quoted price (NGN)'], [
 ('Guide Rice', 'bag', '10', '1000'), ('Guide Flour', 'bag', '10', '1000'),
 ('Guide Syrup', 'can', '5', '500'), ('Guide Tissue', 'pack', '10', '500')])
steps(['Open Master data → Supplier items. Scroll to Link a supplier item.',
       'Select Guide Test Supplier as Vendor and the appropriate GUIDE item as Inventory item.',
       'Enter Purchase unit / pack, Base units per purchase unit and Quoted price per purchase unit (NGN) from the table.',
       'Enter a Pack specification / evidence note, for example: Fictional test bag contains 10 kg. Keep Status Active.',
       'Click Save supplier item. Repeat for all four purchased items. Do not add Guide Sauce here.'])
check('One rice bag equals 10 kg; one syrup can equals 5 l. These mappings and quotes still create no stock or weighted average cost.')

chapter(8, 'Add the sub-recipe and final recipes', 'OWNER with Inventory duty', 'Master data → Sub-recipes / Final recipes')
sub('Create the stocked sub-recipe first')
table(['Field', 'Enter'], [('Recipe name', 'Guide Sauce Batch'), ('Recipe / menu code', 'GUIDE-SUB-SAUCE'),
 ('Semi-finished output SKU', 'Guide Sauce (kg)'), ('Standard output quantity', '4'), ('Item 1', 'Guide Flour (kg)'), ('Quantity 1', '5')])
steps(['Open Master data → Sub-recipes. Scroll to Create a recipe.',
       'Enter the fields above and click Save recipe. The dedicated sub-recipe screen already fixes the recipe type.',
       'This is a standard batch: 5 kg flour produces 4 kg sauce. No stock moves when you save it.'])
sub('Create three final recipes')
table(['Recipe name / code', 'Standard output', 'Ingredient lines'], [
 ('Guide Rice Meal / GUIDE-FOOD-RCP', '1', 'Guide Rice: 2 kg'),
 ('Guide Syrup Drink / GUIDE-DRINK-RCP', '1', 'Guide Syrup: 0.25 l'),
 ('Guide Mixed Plate / GUIDE-MIX-RCP', '1', 'Guide Flour: 0.1 kg; Guide Sauce: 0.2 kg')])
steps(['Open Master data → Final recipes → Create a recipe.',
       'For each row, enter Recipe name, Recipe / menu code and Standard output quantity. Leave Semi-finished output SKU blank for a final recipe.',
       'Choose the ingredient and its Quantity. For Guide Mixed Plate, click Add another item to add the second ingredient.',
       'Click Save recipe for each final recipe. Ingredient quantities are for the entire standard output.'])
note('An Incomplete cost before approved purchasing is expected: item references and quotes are not inventory averages. Do not enter fabricated opening stock just to remove that message. If a recipe is wrong, use Create a new recipe version and enter the full corrected ingredient list.')
check('You have one sub-recipe and three final recipes. No recipe creation changed stock quantities.')

chapter(9, 'Map menu/POS codes to the final recipes', 'OWNER with Inventory duty', 'Master data → Menu mapping')
table(['POS / menu code', 'Menu item name', 'Final recipe', 'Sales category', 'Usual preparation location'], [
 ('GUIDE-FOOD', 'Guide Rice Meal', 'GUIDE-FOOD-RCP', 'Food', 'Kitchen'),
 ('GUIDE-DRINK', 'Guide Syrup Drink', 'GUIDE-DRINK-RCP', 'Beverage', 'Bar'),
 ('GUIDE-MIX', 'Guide Mixed Plate', 'GUIDE-MIX-RCP', 'Food', 'Kitchen')])
steps(['Open Master data → Menu mapping. Scroll to Add menu mapping.',
       'Enter the first row’s POS / menu code and Menu item name.',
       'Choose Sales category, Final recipe and Usual preparation location exactly as shown. Keep Status Active.',
       'Click Save menu mapping. Repeat for the other two rows.'])
note('The sales file uses GUIDE-FOOD, GUIDE-DRINK and GUIDE-MIX, not the recipe codes ending in -RCP. Sales category determines the food/beverage revenue split. The location in each imported sales row controls the stock deduction target.')
check('All three menu codes are active and point to the intended final recipe.')

chapter(10, 'Create and independently approve your first order', 'OWNER creates; OPS approves', 'Purchasing')
table(['Item', 'Purchase quantity', 'Price per purchase unit (NGN)', 'Expected stock when fully received'], [
 ('Guide Rice — bag', '2', '1000', '20 kg'), ('Guide Flour — bag', '1', '1000', '10 kg'),
 ('Guide Syrup — can', '2', '500', '10 l'), ('Guide Tissue — pack', '1', '500', '10 each')])
steps(['As Owner, open Purchasing and click the Guide Test Supplier vendor button under New purchase order.',
       'Enter GUIDE-PO-01 in Reason / reference. Expected delivery date is optional.',
       'Select the first item and enter Purchase quantity and Price per purchase unit using the table. Use Add another item for the remaining three lines.',
       'Click Save purchase order draft. Locate the order by GUIDE-PO-01; its PO number may differ from examples in other guides.',
       'Review the four lines. The order value should be NGN 4500. Click Submit order for approval.',
       'Switch to OPS. Open Purchasing, find the same order and click Approve purchase order. Alternatively use Approvals, expand the order and click Approve.',
       'Return to Owner and refresh. The order should show approved.'])
check('Order value is NGN 4500. All ordered stock quantities are still zero before approved receipts. Owner cannot receive the order they created.')

chapter(11, 'Receive the order in two deliveries', 'STORE receives; OPS approves', 'Purchasing → your approved order')
sub('Delivery 1 — receive only one rice bag')
steps(['As STORE, open Purchasing. Find GUIDE-PO-01 and expand Record accepted delivery.',
       'Set Receiving date to D. Select Main Store as Receiving location. Enter GUIDE-RECEIPT-01 in Reason / reference.',
       'Select Guide Rice (bag) as Item 1. Enter Accepted purchase quantity 1 = 1. Leave Rejected purchase quantity 1 = 0.',
       'Click Save receiving draft. The app opens Documents & approvals. Find receipt GUIDE-RECEIPT-01 and click Submit for approval.',
       'Switch to OPS → Inventory → Documents & approvals. Find that receipt, review it and click Approve and post.',
       'As STORE, refresh Stock on hand: Main Store should have 10 kg rice. Purchasing should still show one rice bag and the other items outstanding.'])
sub('Delivery 2 — receive the remaining order')
table(['Item', 'Accepted purchase quantity', 'Rejected purchase quantity'], [('Guide Rice (bag)', '1', '0'),
    ('Guide Flour (bag)', '1', '0'), ('Guide Syrup (can)', '2', '0'), ('Guide Tissue (pack)', '1', '0')])
steps(['As STORE, reopen the same approved order → Record accepted delivery. Use Receiving date D, Receiving location Main Store and reason GUIDE-RECEIPT-02.',
       'Enter the remaining four lines above using Add another item. These are purchase units, not kg/l/each.',
       'Save receiving draft, then Submit for approval from its receipt card.',
       'As OPS, review and Approve and post the second receipt.'])
check('Main Store: Rice 20 kg, Flour 10 kg, Syrup 10 l, Tissue 10 each. Receipt value totals NGN 4500. Latest base-unit prices: rice/flour/syrup NGN 100; tissue NGN 50. The order has no outstanding quantities.')
note('Rejected quantity is documented separately and never becomes stock. Keep it zero for this baseline. Draft receipts do not change stock until approved and posted.')

chapter(12, 'Request, dispatch and confirm department stock', 'KITCHEN / BAR request; OPS approves; STORE dispatches', 'Stock requests')
sub('Kitchen request')
steps(['As KITCHEN, open Stock requests → Request stock. Requesting location should be Kitchen.',
       'Use reason GUIDE-KITCHEN-REQ. Add Guide Rice 10 kg, Guide Flour 10 kg and Guide Tissue 2 each.',
       'Click Save request draft, then Submit request on that request card.',
       'As OPS, open Stock requests, locate the request and click Approve request. The approval does not reserve or move stock.',
       'As STORE, open the approved request and expand Dispatch approved stock. Review the prefilled quantities, set Dispatch date D and reason GUIDE-KITCHEN-DISPATCH. Click Dispatch entered quantities.',
       'Check that Main Store decreased and a dispatched transfer exists. Kitchen stock has not increased until acknowledgement.',
       'As KITCHEN, open Inventory → Documents & approvals. Find transfer GUIDE-KITCHEN-DISPATCH → Confirm received quantities.',
       'Enter a receipt reason, choose Rice quantity 10, Flour quantity 10 and Tissue quantity 2 with Add another item. Click Confirm receipt. The transfer should become received.'])
sub('Bar request')
steps(['As BAR, repeat the request process for Guide Syrup 5 l and Guide Tissue 2 each. Use GUIDE-BAR-REQ.',
       'OPS approves. STORE dispatches on D with reason GUIDE-BAR-DISPATCH.',
       'BAR confirms the actual Syrup 5 l and Tissue 2 each in Documents & approvals.'])
table(['Item', 'Main Store', 'Kitchen', 'Bar'], [('Rice (kg)', '10', '10', '0'), ('Flour (kg)', '0', '10', '0'),
 ('Syrup (l)', '5', '0', '5'), ('Tissue (each)', '6', '2', '2')])
check('Balances match the table; both transfers are received. Department accounts show only authorized departmental stock. Owner/Store can see the wider inventory scope.')
note('For a partial acknowledgement test, confirm only part first, inspect the remaining transit, then confirm only the remainder. Do not enter the full original quantity again. Unresolved shortage settlement is a separate feature still to complete.')

chapter(13, 'Import and approve the baseline sales', 'OWNER with Inventory duty; OPS approves', 'Inventory → Sales import')
p('Use the included file ProfitPlate_Guide_Baseline_Sales.csv. It contains exactly the two rows in Appendix A. Do not add the later mixed-plate sale to this baseline file.')
steps(['As Owner, open Inventory → Sales import → Import sales.',
       'Select the baseline CSV under Sales CSV file. Wait for 2 sales lines ready for validation.',
       'Check that Food quantity is 3 with net sales 1500 and Kitchen target; Drink quantity is 4 with net sales 1000 and Bar target.',
       'Enter GUIDE-POS as Source system, set Sales business date to D and enter GUIDE-BASELINE-SALES in Import reference / reason.',
       'Click Validate and save sales draft. Review the ingredient deductions in its document card.',
       'Click Submit for approval. As OPS, review the same sales document and click Approve and post.',
       'Refresh Stock on hand and Movements. Guide Rice in Kitchen decreases by 6 kg; Guide Syrup in Bar decreases by 1 l. Other locations do not deplete for these sales.',
       'Try importing the exact same baseline file with source GUIDE-POS again. The already-posted source IDs must be rejected; no second depletion should occur.'])
table(['After baseline sales', 'Main Store', 'Kitchen', 'Bar'], [('Rice (kg)', '10', '4', '0'), ('Flour (kg)', '0', '10', '0'),
 ('Syrup (l)', '5', '0', '4'), ('Tissue (each)', '6', '2', '2')])
check('Rice deduction is 6 kg; syrup deduction is 1 l. Re-importing posted source IDs does not change these balances again.')
note('net_sales means the total revenue for that line after discounts and excluding tax, in NGN. It is not price per unit. Keep source IDs stable and unique. Missing revenue is unavailable; explicit zero means free. This importer rejects refunds and negative quantities.')

chapter(14, 'Check margins, reconcile and publish your first report', 'OWNER prepares; OPS approves reconciliation', 'Reports')
sub('Confirm the theoretical baseline first')
steps(['As Owner, open Reports. Choose Daily. Set Reference / start date to D and choose Whole business as Location. If using End date, set it to D too. Click View report.',
       'Confirm the header dates are D to D. There must be exactly 2 posted imported sales lines, with both revenue amounts supplied.',
       'Check the expected theoretical values in the table. These totals assume no other sales were posted for D in this restaurant.',
       'A late-entry indicator is expected because the fictional transactions were entered today with business date D. Actual COGS should remain Unavailable before reconciliation.'])
table(['Measure', 'Expected baseline'], [('Net sales', 'NGN 2500'), ('Theoretical COGS', 'NGN 700'),
 ('Theoretical gross profit', 'NGN 1800'), ('Theoretical gross margin', '72%'),
 ('Theoretical food cost / net food sales', '40%'), ('Accepted purchase receipt value', 'NGN 4500')])
sub('Demonstrate the controlled Finance workflow')
note('The actual cost values below are fictional test declarations. They test entry, independent approval and formulas. They are not automatically derived or independently verified accounting COGS. For real reports, use an actual accounting/count reconciliation, not these numbers.')
steps(['Ensure every stock record dated D is posted/received or rejected, with no draft/submitted stock documents or counts left for D.',
       'As Owner, scroll to Submit Finance reconciliation. Keep Whole business selected and the report at D to D.',
       'Enter Reconciled POS line count = 2; Reconciled net sales (NGN) = 2500; Actual COGS (NGN) = 900; Actual food cost (NGN) = 650.',
       'Enter Supporting evidence reference = GUIDE-FIN-001 — Appendix B fictional schedule. Enter reason Fictional Finance workflow test. Click Submit reconciliation.',
       'Switch to OPS → Reports and open the same D-to-D range. Find Accounting reconciliations, review the submitted record and click Approve reconciliation.',
       'Confirm Gross profit = NGN 1600, Gross margin = 64%, Food cost = NGN 650 and Food cost / net food sales = 43.3333%.',
       'As Owner, open the same report → Publish a report snapshot. Enter GUIDE-BASELINE-PUBLISHED as Reason / reference and click Publish snapshot.',
       'Under Published reports, click View saved report. Record its snapshot ID and revision. Click Download spreadsheet (CSV) and open the downloaded file in Excel.'])
check('The published report retains the selected dates, revision and figures. Actual figures appeared only after the independent reconciliation approval.')

chapter(15, 'Test a later purchase and weighted average cost', 'OWNER orders; STORE receives; OPS approves', 'Purchasing, then Inventory → Stock on hand')
p('Complete and publish section 14 before doing this test. Use T for the new receipt, so the D baseline remains a separate historical report.')
steps(['As Owner, create a new order from Guide Test Supplier for Guide Rice: Purchase quantity 1 bag and Price per purchase unit 2000. Use reason GUIDE-RICE-COST-CHANGE.',
       'Submit the new order. As OPS, approve it. The approved order explicitly authorizes the new negotiated price.',
       'As STORE, record delivery of that 1 bag into Main Store with Receiving date T and reason GUIDE-NEW-RICE-RECEIPT. Save, submit and have OPS Approve and post.',
       'Check Main Store rice: it should increase from 10 kg to 20 kg. Whole-business rice quantity increases from 14 kg to 24 kg.',
       'Check the latest rice purchase price: NGN 200/kg. The business-wide weighted average should be approximately NGN 141.66666667/kg.',
       'Open the final recipe master: Guide Rice Meal current recipe cost should be approximately NGN 283.33333334 per meal.',
       'Open the saved D baseline snapshot again. Its theoretical COGS remains NGN 700 and its approved actual gross margin remains 64%. A later-dated purchase must not overwrite it.'])
table(['Arithmetic', 'Result'], [('Rice remaining after D sales', '14 kg × NGN 100 = NGN 1400'),
 ('New rice received on T', '10 kg × NGN 200 = NGN 2000'), ('New business-wide rice average', '(1400 + 2000) / (14 + 10) = NGN 141.66666667/kg')])
check('Latest purchase price, current weighted average and historical reported cost remain distinct.')

chapter(16, 'Count stock and approve a controlled adjustment', 'KITCHEN counts; OPS approves', 'Stock counts')
steps(['As KITCHEN, open Stock counts → Record a physical count. Count location should be Kitchen.',
       'For this fictional exercise, enter reason GUIDE-KITCHEN-COUNT. Select Guide Rice and Physical quantity 1 = 3. Kitchen book rice should still be 4 kg before this step.',
       'Pause movements for that counted item. Click Save count draft. Review the saved Book 4, Physical 3 and Variance -1 kg.',
       'Click Submit count for approval. Your Kitchen account must not have an approval action for its own submission.',
       'As OPS, open Stock counts, review GUIDE-KITCHEN-COUNT and click Approve count adjustment.',
       'Refresh Kitchen stock: rice is now 3 kg. Inspect Movements for count adjustment. It should show a single 1 kg deduction; the physical balance was not overwritten without evidence.'])
note('Counts use an exact server cutoff and today’s business date. If the item moves at the counted location after capture, approval is blocked and a new count is required. A positive adjustment needs an established inventory cost. Counting one item does not certify that all items have been counted.')
check('Kitchen rice is 3 kg; the count is posted; repeated clicks/retries do not apply the adjustment twice.')

chapter(17, 'Record waste in all areas and use general supplies', 'KITCHEN, BAR and STORE create; OPS approves', 'Inventory → Documents & approvals')
sub('Waste tests — use T')
table(['Creator', 'Source location', 'Item / quantity', 'Reason'], [('KITCHEN', 'Kitchen', 'Guide Rice / 0.5 kg', 'GUIDE-KITCHEN-WASTE'),
 ('BAR', 'Bar', 'Guide Syrup / 0.5 l', 'GUIDE-BAR-WASTE'), ('STORE', 'Main Store', 'Guide Rice / 0.5 kg', 'GUIDE-STORE-WASTE')])
steps(['Sign in with the creator account for the first row. Open Inventory → Documents & approvals → waste.',
       'Set Event date T, enter the listed Reason / reference and confirm Source location. Choose the listed item and quantity.',
       'Click Save draft, then Submit for approval. As OPS, review that exact reason and click Approve and post.',
       'Repeat for the other two rows. Expected balances: Kitchen Rice 2.5 kg; Main Store Rice 19.5 kg; Bar Syrup 3.5 l.'])
sub('General supply usage — Kitchen tissue')
steps(['As KITCHEN, open Inventory → Documents & approvals → General-item usage.',
       'Set Event date T, Source location Kitchen and reason GUIDE-TISSUE-USAGE. Enter Guide Tissue quantity 1.',
       'Save and submit the draft. OPS reviews and approves it.',
       'Kitchen tissue decreases from 2 each to 1 each. General supplies are recorded through usage, not inserted into food recipes.'])
check('Each waste affects only its own location. Tissue usage is a separate movement type. Do not record the same physical loss as both waste and count shortage.')

chapter(18, 'Produce a semi-finished SKU and sell a mixed recipe', 'KITCHEN produces; OWNER imports sales; OPS approves', 'Inventory → Documents & approvals → production')
sub('Production with a lower actual yield')
steps(['As KITCHEN, choose production. Set Event date T, Source location Kitchen and reason GUIDE-SAUCE-PRODUCTION.',
       'Choose the saved Guide Sauce Batch production recipe version. Check its standard output is 4.',
       'Enter Planned output quantity = 4 and Actual declared output quantity = 3.6.',
       'Click Save draft. Review the standard input deduction of Guide Flour 5 kg, then Submit for approval.',
       'OPS reviews and approves/posts. Kitchen Flour decreases from 10 to 5 kg. Kitchen Sauce increases from zero to 3.6 kg.',
       'The batch input value is NGN 500. Its output carrying cost is approximately NGN 138.88888889/kg. The live standard sauce recipe cost remains NGN 125/kg at flour cost NGN 100/kg. The difference reflects the lower actual yield.'])
sub('Optional mixed-plate sale')
steps(['As Owner, import the included ProfitPlate_Guide_Mixed_Plate_Sales.csv using Source system GUIDE-POS and Sales business date T. Use reason GUIDE-MIXED-SALE.',
       'Save and submit. OPS approves it. This is a new source ID; it does not replace the baseline file.',
       'One mixed plate deducts Guide Flour 0.1 kg and stocked Guide Sauce 0.2 kg. Kitchen Flour becomes 4.9 kg and Sauce becomes 3.4 kg.',
       'The sauce’s underlying flour is not deducted a second time. The mixed recipe’s theoretical cost is NGN 35; the stock carrying-value deduction is approximately NGN 37.77777778.',
       'A T-only report can show the mixed plate’s theoretical margin, but actual reconciliation for T must wait until the business date is completed.'])
check('Production uses standard inputs scaled to planned output, adds actual output, and final sales consume the semi-finished stock once.')

chapter(19, 'Test safe failures, periods and account restrictions', 'Use the account named in each test')
table(['Test', 'How to try it', 'Expected safe result'], [
 ('Duplicate sales', 'Owner re-imports the baseline source IDs using GUIDE-POS.', 'Rejected; no repeated stock deduction.'),
 ('No self-approval', 'Review your own submitted document; for a manager-owned test use OPS plus a separate backup approver.', 'Creator cannot approve. Another eligible Operations Manager is required.'),
 ('Over-receipt', 'STORE attempts more than the approved order’s outstanding quantity.', 'Rejected; no extra stock posted.'),
 ('Wrong department', 'BAR attempts to request Kitchen stock or open its transaction.', 'Access/action denied; department scope preserved.'),
 ('Insufficient stock', 'KITCHEN saves/submits a clearly identified test waste above its available quantity; OPS attempts posting.', 'Posting fails atomically. OPS should reject that test submission afterwards.'),
 ('Closed period', 'OPS closes the period after all positive exercises. Another account submits a dated test document; OPS attempts posting.', 'Posting blocked. OPS can reopen with a reason; reject unused test documents.'),
 ('Suspended account', 'Owner suspends a non-owner test account from Team and access → Manage this member.', 'An already-open form cannot save. Reactivate after the test.'),
 ('Admin separation', 'ADMIN opens Reports or tries to create a receipt without extra duties.', 'No financial report/posting access by default.')])
sub('Optional precise stale-count test')
steps(['After section 17, Kitchen Rice should be 2.5 kg. KITCHEN saves and submits a count of 2.4 kg with reason GUIDE-STALE-COUNT.',
       'Before approving that count, KITCHEN records a separate, fictional waste of 0.1 kg on T with reason GUIDE-AFTER-COUNT-WASTE. OPS approves the waste.',
       'OPS now attempts Approve count adjustment on GUIDE-STALE-COUNT. It must be blocked because the counted stock moved after the cutoff.',
       'OPS rejects the stale count with a clear reason. KITCHEN creates a fresh count of the new physical 2.4 kg and submits it; OPS approves. With Book 2.4 and Physical 2.4, no extra quantity adjustment is needed.'])
check('Blocked actions do not partially post. Reject intentionally failed submissions so they do not remain pending in later reports.')

chapter(20, 'Review reports and assign a follow-up action', 'OWNER or OPS; assignee updates their own action', 'Reports')
steps(['Open Reports. Choose Daily and T to review today’s advanced tests, or D for the baseline. Confirm exact dates and location before interpreting amounts.',
       'Choose Weekly to see a Monday–Sunday period, or Monthly for the calendar month. Click View report and read the displayed boundaries. A custom End date overrides the automatic end date.',
       'Review Purchasing and losses, Inventory at the reporting cutoff and Movement evidence. Expand movement rows and use View source document links.',
       'On a live report, read Previous-period comparison. Missing previous-period evidence should be Unavailable, not zero profit.',
       'To assign an action, expand Live management actions → Create management action. Enter Action title Review test rice shortage, choose KITCHEN as Action owner, set Priority High and a Due date.',
       'Enter Action evidence / finding reference = GUIDE-KITCHEN-COUNT, then click Create action.',
       'As KITCHEN, open Reports → Live management actions. Expand Update progress, choose In progress and enter a reason; click Save action progress. Later resolve it with an explanatory note.'])
note('Management actions remain live even while viewing a saved report. Progress notes do not post inventory or amend published financial figures. The current app does not automatically send report emails, publish PDF reports or operate a platform-admin console.')
check('The assignee can update their own action, department reports hide company financial cards, and old published snapshots remain unchanged.')

chapter(21, 'Troubleshooting before you repeat a transaction')
table(['What you see', 'What to check / do'], [
 ('A button is missing', 'Check the signed-in email, working role, explicit duties and current status. Owner needs Procurement/Inventory duties for those tasks. Expand the record and scroll down.'),
 ('Cannot receive your order', 'The purchase order creator cannot receive it. Use STORE, not the Owner who created GUIDE-PO-01.'),
 ('No Approve action', 'Use a separately appointed OPS account and a submitted record created by a different user. Owner status alone is not approval authority.'),
 ('Posting period error', 'Check the record’s business date and an open, non-overlapping period. Transfer confirmation also needs today in an open period.'),
 ('Incomplete recipe cost', 'Check approved receipts and ingredient costs. Quoted or imported reference prices do not establish inventory averages.'),
 ('Margin is Unavailable', 'Check net_sales values, posted sales status, recipe-cost evidence and food/beverage classifications. Actual figures need an approved matching reconciliation.'),
 ('Reconciliation mismatch', 'Use a completed D-to-D range and Whole business. Match the report’s posted source-line count/net revenue; clear pending stock documents/counts for that period.'),
 ('Stale reconciliation', 'Source evidence changed after preparation. Reject the stale submission and prepare a fresh supported reconciliation. Published snapshots remain intact.'),
 ('CSV rejected', 'Use exactly the supplied headers, plain decimal numbers, active menu codes and unique source IDs. Do not type currency symbols or commas inside unquoted amounts.'),
 ('Count approval blocked', 'Stock moved after capture. Reject the stale count, recount and submit a new record.'),
 ('Phone cannot connect', 'Check the working phone URL, server and network first. Do not use phone localhost for the laptop server.'),
 ('Account service limiting requests', 'Wait and follow the displayed recovery screen. Avoid repeated sign-in clicks. Capture the error time and page.'),
 ('Save result uncertain', 'Refresh and check for the reason/source ID before starting another record. Retry with the same unchanged details if appropriate.'),
 ('A posted record is wrong', 'Record its reference and stop that test. Do not delete stock history or invent offsetting waste. Approved reversal/return workflows are still pending.')])

chapter(22, 'Your acceptance-test record')
p('Use the included CSV test log or fill this table. Record references and screenshots without passwords or invitation codes.')
table(['ID', 'Test', 'Pass / Fail', 'Reference / notes'], [
 ('01', 'Owner login and refresh', '', ''), ('02', 'Invitations and correct memberships', '', ''),
 ('03', 'Owner duties and separate OPS appointment', '', ''), ('04', 'Locations and open periods', '', ''),
 ('05', 'Five items / vendor / packs', '', ''), ('06', 'Sub-recipe / final recipes / menu mappings', '', ''),
 ('07', 'Order value 4500; approval creates no stock', '', ''), ('08', 'Partial and final receiving balances', '', ''),
 ('09', 'Kitchen and Bar request / dispatch / receipt', '', ''), ('10', 'Baseline sales and duplicate protection', '', ''),
 ('11', 'Theoretical COGS 700 / margin 72%', '', ''), ('12', 'Reconciled COGS 900 / margin 64%', '', ''),
 ('13', 'Published snapshot / CSV export', '', ''), ('14', 'New rice average 141.66666667', '', ''),
 ('15', 'Count adjustment and stale-count block', '', ''), ('16', 'Waste at three areas / tissue usage', '', ''),
 ('17', 'Production and single semi-finished depletion', '', ''), ('18', 'Role / suspension / closed-period checks', '', ''),
 ('19', 'Management action ownership', '', ''), ('20', 'Phone and desktop walkthrough', '', '')])
sub('When you report an issue')
bullets(['Say which account role was signed in, which screen you opened and which button you clicked.',
         'Include the record reason/reference, selected dates and expected versus actual result.',
         'Describe whether it happens on laptop, phone or both. Include the exact message or a screenshot with credentials hidden.',
         'Do not reset or delete the database to repeat a test. Preserve the evidence so the cause can be investigated.'])

chapter(23, 'Plain-English glossary and next steps')
table(['Term', 'Meaning in this guide'], [('Owner', 'Your restaurant’s oversight and team account, not the global ProfitPlate platform administrator.'),
 ('Master data', 'Reusable definitions such as items, suppliers, recipes, units and menu codes.'),
 ('SKU', 'A stable code identifying one inventory item.'), ('Base unit', 'The unit used to count stock, for example kg, l or each.'),
 ('Purchase pack', 'The supplier’s buying unit; one bag may contain 10 kg.'), ('Draft', 'Saved evidence that has not yet been submitted or posted.'),
 ('Submitted', 'Waiting for an eligible independent approver.'), ('Posted', 'Approved and applied to the inventory ledger.'),
 ('Dispatched / in transit', 'Stock has left Main Store but the destination has not confirmed all of it.'),
 ('Received', 'The transfer’s quantities have been acknowledged at destination.'), ('Weighted average cost', 'The business-wide average inventory cost per item after valued stock changes.'),
 ('Theoretical COGS', 'Standard recipe ingredient cost for the imported sales quantities.'),
 ('Actual COGS in this version', 'Finance-supplied accounting cost with supporting evidence and independent Operations approval.'),
 ('Snapshot', 'A saved report version whose original figures do not change.'), ('Reconciliation', 'Checking the report against its POS and accounting evidence.'),
 ('Cutoff', 'The date or exact instant used as the reference for a report or count.')])
sub('When the walkthrough is complete')
steps(['Collect your test log and screenshots. Identify any failed or confusing steps before adding more transactions.',
       'Keep the GUIDE records identified as test evidence. Do not treat their fictional actual-cost declarations as real restaurant performance.',
       'Agree a controlled opening-stock plan and reviewed real masters before live operation.',
       'Next product-readiness work includes approved corrections/reversals, returns/refunds and transfer-shortage settlement; production hosting/recovery/security checks; and a hosted operator pilot.',
       'An overall ProfitPlate Platform Console and broader multi-restaurant owner onboarding are separate work. Do not give customer restaurants direct Supabase project access.'])

chapter('APPENDIX A', 'Ready-to-use sales CSV files')
p('The downloadable starter pack contains these CSV files. Keep the headers exactly as supplied. If using Excel to edit them, save as CSV UTF-8, not an Excel workbook renamed .csv.')
sub('ProfitPlate_Guide_Baseline_Sales.csv — import on D')
baseline = 'external_id,menu_code,quantity,location,net_sales\nGUIDE-POS-001,GUIDE-FOOD,3,KITCHEN,1500\nGUIDE-POS-002,GUIDE-DRINK,4,BAR,1000\n'
for line in baseline.strip().splitlines(): p(line)
sub('ProfitPlate_Guide_Mixed_Plate_Sales.csv — optional import on T')
mixed = 'external_id,menu_code,quantity,location,net_sales\nGUIDE-POS-MIX-001,GUIDE-MIX,1,KITCHEN,200\n'
for line in mixed.strip().splitlines(): p(line)
note('Use Source system GUIDE-POS for both files. If repeating the same scenario in a deliberately fresh test run, create a new unique prefix and update all related masters/source IDs consistently. Reusing posted source IDs is expected to fail.')
sub('Opening the starter pack')
steps(['Download the ZIP on your laptop and extract it to a normal folder.',
       'Open this Word guide. Keep both sales CSV files unchanged for the first walkthrough.',
       'Use the CSV test log to record your results. It is not an inventory import file.'])

chapter('APPENDIX B', 'Arithmetic and the fictional Finance schedule')
sub('Baseline source arithmetic')
table(['Source', 'Calculation', 'Result'], [('Rice received', '2 bags × 10 kg × NGN 100/kg', '20 kg / NGN 2000'),
 ('Flour received', '1 bag × 10 kg × NGN 100/kg', '10 kg / NGN 1000'),
 ('Syrup received', '2 cans × 5 l × NGN 100/l', '10 l / NGN 1000'),
 ('Tissue received', '1 pack × 10 each × NGN 50/each', '10 each / NGN 500'),
 ('Total purchase receipts', '2000 + 1000 + 1000 + 500', 'NGN 4500'),
 ('Food theoretical cost', '3 meals × 2 kg × NGN 100/kg', 'NGN 600'),
 ('Drink theoretical cost', '4 drinks × 0.25 l × NGN 100/l', 'NGN 100'),
 ('Net sales', '1500 food + 1000 beverage', 'NGN 2500'),
 ('Theoretical COGS', '600 + 100', 'NGN 700'),
 ('Theoretical gross profit / margin', '(2500 − 700) / 2500', 'NGN 1800 / 72%'),
 ('Theoretical food cost %', '600 / 1500 × 100', '40%')])
sub('GUIDE-FIN-001 — fictional schedule for the approval test only')
table(['Declared test evidence', 'Value'], [('Business period', 'D to D — write your actual selected D date'),
 ('Posted POS source-line count', '2'), ('Reconciled net sales', 'NGN 2500'),
 ('Declared actual food cost', 'NGN 650'), ('Declared other/beverage cost for this example', 'NGN 250'),
 ('Declared actual COGS total', 'NGN 900'), ('Resulting gross profit', 'NGN 1600'),
 ('Resulting gross margin', '64%'), ('Resulting actual food cost %', '43.3333%')])
p('Prepared by: ____________________   Reviewed by independent OPS: ____________________')
note('These declared test costs are not a calculation of physical consumption from the test ledger. This schedule demonstrates that the app requires a matching source period and separate approval before displaying actual financial measures. Real accounting evidence must explain its own cost basis and differences.')

chapter('APPENDIX C', 'Checkpoints and finish line')
table(['Checkpoint', 'Expected result'], [('After full initial receiving', 'Main Store: Rice 20 kg; Flour 10 kg; Syrup 10 l; Tissue 10 each.'),
 ('After department receipts', 'Main: Rice 10, Syrup 5, Tissue 6. Kitchen: Rice 10, Flour 10, Tissue 2. Bar: Syrup 5, Tissue 2.'),
 ('After baseline sales', 'Kitchen Rice 4 kg; Bar Syrup 4 l. Theoretical COGS 700; theoretical margin 72%.'),
 ('After baseline Finance approval', 'Reconciled COGS 900; gross profit 1600; gross margin 64%; actual food cost 43.3333%.'),
 ('After later rice purchase, before count/waste', 'Rice across business 24 kg at approximately NGN 141.66666667/kg.'),
 ('After count and three waste tests', 'Kitchen Rice 2.5 kg; Main Rice 19.5 kg; Bar Syrup 3.5 l.'),
 ('After tissue usage', 'Tissue: Main 6 each, Kitchen 1 each, Bar 2 each.'),
 ('After production and optional mixed sale', 'Kitchen Flour 4.9 kg; Kitchen Sauce 3.4 kg.'),
 ('After optional stale-count exercise', 'Kitchen Rice 2.4 kg. Old stale count rejected; fresh equal count posted without another deduction.')])
sub('Final checklist')
for text in ['All expected positive tests passed or have a documented issue.',
             'Failed test submissions were rejected where appropriate; no unexplained pending records remain.',
             'The baseline snapshot and exported CSV are saved with the test log.',
             'No passwords or invitation codes appear in shared screenshots or documents.',
             'The same account was never used to approve its own transaction.',
             'You can explain the difference between reference price, latest purchase price, weighted average and historical report cost.',
             'You understand that actual COGS is supported Finance evidence in this version.',
             'Production release is a separate decision after corrections, recovery and real operator testing.']:
    p('[ ] ' + text)

# Independent arithmetic checks guard the handbook's reference values.
assert D('20') * D('100') + D('10') * D('100') + D('10') * D('100') + D('10') * D('50') == D('4500')
assert D('3') * D('2') * D('100') + D('4') * D('.25') * D('100') == D('700')
assert (D('2500') - D('700')) / D('2500') * 100 == D('72')
assert (D('2500') - D('900')) / D('2500') * 100 == D('64')
assert (D('650') / D('1500') * 100).quantize(D('.0001')) == D('43.3333')
assert ((D('1400') + D('2000')) / D('24')).quantize(D('.00000001')) == D('141.66666667')
assert D('10') - D('5') - D('.1') == D('4.9')
assert D('3.6') - D('.2') == D('3.4')
path = OUT / f'{NAME}.docx'
doc.save(path)
(OUT / 'ProfitPlate_Guide_Baseline_Sales.csv').write_text(baseline, encoding='utf-8-sig', newline='')
(OUT / 'ProfitPlate_Guide_Mixed_Plate_Sales.csv').write_text(mixed, encoding='utf-8-sig', newline='')
log = io.StringIO(newline='')
w = csv.writer(log)
w.writerow(['Test_ID', 'Test_name', 'Result_Pass_Fail', 'Date', 'Role_used', 'Record_reference', 'Expected_result', 'Actual_result', 'Notes'])
for i, title in enumerate(['Owner login', 'Invitations', 'Owner duties and OPS', 'Locations and period', 'Masters', 'Recipes and menu mappings', 'Purchase approval', 'Partial receiving', 'Department transfers', 'Sales and duplicates', 'Theoretical report', 'Finance reconciliation', 'Publication and export', 'Weighted average', 'Counts and stale cutoff', 'Waste and usage', 'Production and mixed sale', 'Permissions and period locks', 'Management actions', 'Phone and desktop'], 1):
    w.writerow([f'{i:02}', title, '', '', '', '', '', '', ''])
(OUT / 'ProfitPlate_Guide_Test_Log.csv').write_text(log.getvalue(), encoding='utf-8-sig', newline='')
files = [path, OUT / 'ProfitPlate_Guide_Baseline_Sales.csv', OUT / 'ProfitPlate_Guide_Mixed_Plate_Sales.csv', OUT / 'ProfitPlate_Guide_Test_Log.csv']
pack = OUT / 'ProfitPlate_Owner_Guide_Starter_Pack.zip'
with zipfile.ZipFile(pack, 'w', compression=zipfile.ZIP_DEFLATED) as z:
    for f in files: z.write(f, f.name)
# Validate the OOXML archive, reopen it, and check complete major sections and templates.
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    assert 'word/document.xml' in z.namelist()
check_doc = Document(path)
headings = [x.text for x in check_doc.paragraphs if x.style.name == 'Heading 1']
assert len(headings) == 27, (len(headings), headings)
assert any('Invite each account' in x for x in headings)
assert any('Check margins' in x for x in headings)
assert any('Troubleshooting' in x for x in headings)
for text, expected in [(baseline, 2), (mixed, 1)]:
    data = list(csv.DictReader(io.StringIO(text)))
    assert len(data) == expected
    assert set(data[0]) == {'external_id', 'menu_code', 'quantity', 'location', 'net_sales'}
print(f'Created and validated: {path}')
print(f'Chapters: {len(headings)}; tables: {len(check_doc.tables)}; paragraphs: {len(check_doc.paragraphs)}')
print(f'Starter pack: {pack}')
for f in files: print(f'{f.name}: {f.stat().st_size:,} bytes')
